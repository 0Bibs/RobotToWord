#!/usr/bin/env python3
"""Converte relatorios em PDF (Robot Structural Analysis) em documentos Word.

Cada pagina do PDF vira um "print" (imagem) dentro do .docx, recortado nas
margens brancas e ajustado a largura util da pagina - o mesmo protocolo feito
manualmente com a ferramenta de captura.

Uso basico:

    python robot_to_word.py "Ligacao 01 - Robot.pdf"
    python robot_to_word.py pasta_com_pdfs/ --out-dir saida/
    python robot_to_word.py *.pdf --merge "Memorial de Ligacoes.docx"
"""

from __future__ import annotations

import argparse
import io
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator

import fitz  # PyMuPDF
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, Twips

# Geometria copiada do modelo "Exemplo de forma para prints.docx":
# A4 retrato, margens de 3,0 cm nas laterais e 2,5 cm em cima/embaixo,
# o que da 15,0 cm de largura util para os prints.
PAGE_WIDTH = Twips(11906)
PAGE_HEIGHT = Twips(16838)
MARGIN_SIDE = Twips(1701)
MARGIN_TOP_BOTTOM = Twips(1417)


PROBE_DPI = 50  # resolucao barata usada so para descobrir onde esta o conteudo


@dataclass(frozen=True)
class Options:
    dpi: int
    crop_mode: str
    crop_threshold: int
    crop_padding_pt: float
    image_format: str
    jpeg_quality: int
    page_break: bool
    title: bool


def render_page(page: "fitz.Page", dpi: int) -> Image.Image:
    """Rasteriza uma pagina do PDF na resolucao pedida."""
    pixmap = page.get_pixmap(dpi=dpi, alpha=False)
    return Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)


def content_bbox(image: Image.Image, threshold: int) -> tuple[int, int, int, int] | None:
    """Retangulo que envolve o conteudo, ignorando o fundo branco.

    Retorna None quando a pagina esta em branco.
    """
    # Tudo acima do limiar vira 0 (fundo); o conteudo vira 255, e getbbox()
    # devolve o retangulo que envolve os pixels nao nulos.
    mask = image.convert("L").point(lambda value: 0 if value >= threshold else 255)
    return mask.getbbox()


def document_columns(pdf: "fitz.Document", threshold: int) -> tuple[float, float] | None:
    """Faixa horizontal ocupada pelo conteudo em todo o PDF, como fracoes 0..1.

    Serve para recortar todas as paginas na mesma largura, mantendo o texto no
    mesmo tamanho de uma pagina para outra. Passada rapida, em baixa resolucao.
    """
    left = right = None
    for page in pdf:
        image = render_page(page, PROBE_DPI)
        bbox = content_bbox(image, threshold)
        if bbox is None:
            continue
        page_width = image.width
        left = bbox[0] / page_width if left is None else min(left, bbox[0] / page_width)
        right = bbox[2] / page_width if right is None else max(right, bbox[2] / page_width)

    return None if left is None else (left, right)


def crop_page(image: Image.Image, options: Options, padding_px: int,
              columns: tuple[float, float] | None) -> Image.Image | None:
    """Recorta uma pagina ja rasterizada conforme o modo escolhido."""
    if options.crop_mode == "none":
        return image

    bbox = content_bbox(image, options.crop_threshold)
    if bbox is None:
        return None  # pagina em branco

    width, height = image.size
    left, top, right, bottom = bbox
    if columns is not None:  # modo uniforme: laterais iguais em todas as paginas
        left, right = columns[0] * width, columns[1] * width

    return image.crop(
        (
            max(0, round(left) - padding_px),
            max(0, top - padding_px),
            min(width, round(right) + padding_px),
            min(height, bottom + padding_px),
        )
    )


def encode(image: Image.Image, options: Options) -> io.BytesIO:
    """Serializa a imagem no formato escolhido, pronta para entrar no Word."""
    buffer = io.BytesIO()
    if options.image_format == "jpeg":
        image.save(buffer, format="JPEG", quality=options.jpeg_quality, optimize=True)
    else:
        image.save(buffer, format="PNG", optimize=True)
    buffer.seek(0)
    return buffer


def iter_prints(pdf_path: Path, options: Options) -> Iterator[tuple[io.BytesIO, int, int]]:
    """Gera (imagem, largura_px, altura_px) para cada pagina util do PDF."""
    padding_px = round(options.crop_padding_pt * options.dpi / 72)
    with fitz.open(pdf_path) as pdf:
        columns = (document_columns(pdf, options.crop_threshold)
                   if options.crop_mode == "uniform" else None)
        for page in pdf:
            image = crop_page(render_page(page, options.dpi), options, padding_px, columns)
            if image is None:
                continue  # pagina em branco
            yield encode(image, options), image.width, image.height


def new_document(template: Path | None) -> Document:
    """Cria o documento base, do zero ou a partir de um modelo .docx/.dotx."""
    if template is None:
        document = Document()
        section = document.sections[0]
        section.page_width = PAGE_WIDTH
        section.page_height = PAGE_HEIGHT
        section.left_margin = MARGIN_SIDE
        section.right_margin = MARGIN_SIDE
        section.top_margin = MARGIN_TOP_BOTTOM
        section.bottom_margin = MARGIN_TOP_BOTTOM
        return document

    # Reaproveita estilos, cabecalho/rodape e margens do modelo, mas esvazia o
    # corpo - preservando o <w:sectPr> final, que guarda a configuracao da pagina.
    document = Document(str(template))
    body = document.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return document


def usable_area(document: Document) -> tuple[int, int]:
    """Largura e altura uteis da pagina, em EMU."""
    section = document.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    height = section.page_height - section.top_margin - section.bottom_margin
    return width, height


def add_print(document: Document, stream: io.BytesIO, px_w: int, px_h: int,
              *, page_break: bool) -> None:
    """Insere um print ocupando toda a largura util, sem distorcer a imagem."""
    max_width, max_height = usable_area(document)

    width = max_width
    height = round(width * px_h / px_w)
    if height > max_height:  # print muito alto: limita pela altura da pagina
        height = max_height
        width = round(height * px_w / px_h)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run()
    if page_break:
        # A quebra vai no inicio do proprio paragrafo da imagem: garante um
        # print por pagina sem criar paragrafos vazios sobrando.
        run.add_break(WD_BREAK.PAGE)
    run.add_picture(stream, width=Emu(width), height=Emu(height))


def build(pdf_paths: Iterable[Path], document: Document, options: Options) -> int:
    """Despeja os prints de todos os PDFs no documento. Retorna quantos prints."""
    total = 0
    for pdf_path in pdf_paths:
        if options.title:
            heading = document.add_heading(pdf_path.stem, level=1)
            if total and options.page_break:
                heading.runs[0].add_break(WD_BREAK.PAGE)

        for index, (stream, px_w, px_h) in enumerate(iter_prints(pdf_path, options)):
            first_of_document = total == 0
            # Sem titulo, o primeiro print de cada PDF tambem comeca em pagina nova.
            skip_break = first_of_document or (options.title and index == 0)
            add_print(document, stream, px_w, px_h,
                      page_break=options.page_break and not skip_break)
            total += 1

        print(f"  {pdf_path.name}: {total} print(s) acumulado(s)", file=sys.stderr)
    return total


def collect_pdfs(inputs: Iterable[str]) -> list[Path]:
    """Expande arquivos e pastas em uma lista ordenada de PDFs."""
    pdfs: list[Path] = []
    for raw in inputs:
        path = Path(raw)
        if path.is_dir():
            pdfs.extend(sorted(path.glob("*.pdf")))
        elif path.suffix.lower() == ".pdf":
            pdfs.append(path)
        else:
            raise SystemExit(f"Nao e um PDF nem uma pasta: {path}")

    if not pdfs:
        raise SystemExit("Nenhum PDF encontrado.")

    missing = [p for p in pdfs if not p.is_file()]
    if missing:
        raise SystemExit("Arquivo(s) inexistente(s): " + ", ".join(map(str, missing)))
    return pdfs


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Transfere os prints de relatorios PDF do Robot para dentro do Word.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("inputs", nargs="+", metavar="PDF",
                        help="arquivos .pdf e/ou pastas contendo .pdf")
    parser.add_argument("--merge", metavar="ARQUIVO.docx",
                        help="junta todos os PDFs em um unico Word")
    parser.add_argument("--out-dir", metavar="PASTA",
                        help="pasta de saida (padrao: ao lado de cada PDF)")
    parser.add_argument("--template", metavar="MODELO.docx",
                        help="usa um .docx existente como modelo de estilos e margens")
    parser.add_argument("--dpi", type=int, default=200,
                        help="resolucao de captura das paginas")
    parser.add_argument("--crop-mode", choices=("uniform", "tight", "none"), default="uniform",
                        help="uniform: mesma largura em todas as paginas; "
                             "tight: recorte colado ao conteudo de cada pagina; "
                             "none: mantem as margens brancas do PDF")
    parser.add_argument("--crop-threshold", type=int, default=245, metavar="0-255",
                        help="quao claro um pixel precisa ser para contar como fundo")
    parser.add_argument("--crop-padding", type=float, default=4.0, metavar="PT",
                        help="folga deixada ao redor do conteudo recortado")
    parser.add_argument("--format", dest="image_format", choices=("png", "jpeg"),
                        default="png", help="formato dos prints (jpeg gera arquivo menor)")
    parser.add_argument("--quality", type=int, default=85, metavar="1-95",
                        help="qualidade do JPEG, quando --format jpeg")
    parser.add_argument("--no-page-break", action="store_true",
                        help="deixa os prints fluirem em vez de um por pagina")
    parser.add_argument("--title", action="store_true",
                        help="escreve o nome do arquivo como titulo antes dos prints")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    pdfs = collect_pdfs(args.inputs)

    options = Options(
        dpi=args.dpi,
        crop_mode=args.crop_mode,
        crop_threshold=args.crop_threshold,
        crop_padding_pt=args.crop_padding,
        image_format=args.image_format,
        jpeg_quality=args.quality,
        page_break=not args.no_page_break,
        title=args.title,
    )
    template = Path(args.template) if args.template else None
    out_dir = Path(args.out_dir) if args.out_dir else None
    if out_dir:
        out_dir.mkdir(parents=True, exist_ok=True)

    if args.merge:
        groups = [(Path(args.merge), pdfs)]
    else:
        groups = [((out_dir or pdf.parent) / f"{pdf.stem}.docx", [pdf]) for pdf in pdfs]

    for output, sources in groups:
        if out_dir and args.merge:
            output = out_dir / output.name
        print(f"-> {output}", file=sys.stderr)

        document = new_document(template)
        total = build(sources, document, options)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output))
        print(f"   {total} print(s) gravado(s)\n", file=sys.stderr)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
